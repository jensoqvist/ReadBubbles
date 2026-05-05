import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.styles import Alignment, Color, Fill, Font, PatternFill, Border
from openpyxl.styles.borders import Border, Side
from openpyxl.formatting import Rule
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.drawing.spreadsheet_drawing import TwoCellAnchor, AbsoluteAnchor
from openpyxl.drawing.xdr import XDRPoint2D, XDRPositiveSize2D
from openpyxl.utils.units import pixels_to_EMU, cm_to_EMU
from settings import Settings


class XlChangeNotes():
    """
    Class that creates and formats the change notes sheet
    """
    def __init__(self, wbook, settings = Settings()) -> None:
        self.settings = settings
        self.wbook = wbook
        self.SHEET_NAME = "Change_Notes"
        self.TABLE_NAME = "changeNotesTable"
        self.START_ROW = 5
        self.START_COL = 2
        self.START_COL_LETTER = get_column_letter(self.START_COL)

        self.cols = [
            "Date",
            "Rev",
            "Position Number",
            "Change Type",
            "Description",
            "SSSID",
            "Verified by ID",
            "Verified Date",
            "Comment"            
        ]

        self.END_ROW = self.START_ROW + 1000
        self.END_COL = len(self.cols) + self.START_COL - 1
        self.END_COL_LETTER = get_column_letter(self.END_COL)

        if self.SHEET_NAME not in self.wbook.sheetnames:
            self._create_change_sheet()
            self.sheet = self.wbook[self.SHEET_NAME]
            self.sheet['B2'].value = "CHANGE NOTES"
            self.sheet['B2'].font = Font(bold= True)
            for i, col in enumerate(self.cols):
                self.sheet[f"{get_column_letter(i + self.START_COL)}{self.START_ROW}"].value = col
            self._add_table()
            self._set_format(col="Date")
            self._table_colors()
            self._add_borders()
            self._alignment()
            self._adjust_sizes()
            self._data_validation()
            self._add_responsibilitys()
            self._set_format(col="Position Number")
            for c in [
                    "Description"
                    ]:
                self._wrap_text(c)
        self.wbook[self.SHEET_NAME].sheet_state = 'visible'

    def _create_change_sheet(self):
        self.wbook.create_sheet(self.SHEET_NAME, index= -1)
        
    def _add_table(self):
        start, end= f"{self.START_COL_LETTER}{self.START_ROW}", f"{self.END_COL_LETTER}{self.END_ROW}"
        self.table = Table(displayName= self.TABLE_NAME, ref=f'{start}:{end}')
        style = TableStyleInfo(name="TableStyleMedium16", showFirstColumn=False,
                       showLastColumn=False, showRowStripes=True, showColumnStripes=False)
        self.table.tableStyleInfo = style
        self.sheet.add_table(self.table)
        self._table_colors()

    def _table_colors(self):
        color = self.settings["Colors"]["Scania Blue"]
        for col in range(self.START_COL, self.END_COL + 1):
            cell = self.sheet[f"{openpyxl.utils.get_column_letter(col)}{self.START_ROW}"]
            cell.font = Font(color= color["Font Color"])
            cell.fill = PatternFill(fgColor= color["Fill Color"], fill_type= 'solid')

    def _wrap_text(self, col):
        col_index = self.cols.index(col) + self.START_COL
        for r in range(self.START_ROW + 1, self.sheet.max_row + 1):
            self.sheet.cell(r, col_index).alignment = Alignment(wrapText= True)

    def _adjust_sizes(self):   
        self.sheet.column_dimensions["A"].width = 4
        self.sheet.row_dimensions[self.START_ROW].height = 30
        for i, col in enumerate(self.cols):
            col_index = i +  self.START_COL
            try:
                self.sheet.column_dimensions[openpyxl.utils.get_column_letter(col_index)].width = self.settings["Column Size"][col]
            except:
                self.sheet.column_dimensions[openpyxl.utils.get_column_letter(col_index)].width = 24
          
    def _alignment(self):
        for r in range(self.START_COL, self.END_COL + 1):
            self.sheet.cell(self.START_ROW, r).alignment = Alignment(vertical='top', horizontal='center')
        align = {
            "Date": "right",
            "Position Number": "right",
            "Change Type": "center",
            "SSSID": "center",
            "Verified by ID": "center"
        }
        for k,v in align.items():
            for r in range(self.START_ROW+ 1, self.sheet.max_row + 1):
                self.sheet.cell(r, self.START_COL + self.cols.index(k)).alignment = Alignment(horizontal=v)

    def _add_borders(self):
        thin = Side(style= 'thin')
        thick = Side(style= 'thick')
        for row in range(self.START_ROW, self.END_ROW + 1):
            for col in range(self.START_COL, self.END_COL + 1):
                left, right, top, bottom = thin, thin, thin, thin
                if row == self.START_ROW:
                    top = thick
                if row == self.END_ROW:
                    bottom = thick
                if col == self.START_COL:
                    left = thick
                if col == self.END_COL:
                    right = thick
                self.sheet.cell(row= row, column= col).border = Border(left= left, right= right, top= top, bottom= bottom)

    def _data_validation(self):
        for key, value in self.settings["Change Notes Validation"].items():
            validation_string = ", ".join(value)
            col_index = self.cols.index(key) + self.START_COL
            dv = DataValidation(type= "list", formula1= f'"{validation_string}"', allow_blank= True)
            dv.error = "Invalid Entry"
            dv.errorTitle = "Invalid Entry"
            self.sheet.add_data_validation(dv)
            for r in range(self.START_ROW + 1, self.sheet.max_row + 1):
                dv.add(openpyxl.utils.get_column_letter(col_index) + str(r)) 


    def _add_responsibilitys(self):
        row = self.START_ROW - 1
        responsibilitys = self.settings["Change Responsible"]
        for key, value in responsibilitys.items():
            self.sheet[openpyxl.utils.get_column_letter(self.cols.index(value["Responsibilitys"][0]) + self.START_COL) + str(row)].value = f"Responsible: {key}"
            for resbonsibility in value["Responsibilitys"]:
                color = self.settings["Colors"][value["Color"]]
                col_index = self.cols.index(resbonsibility) + self.START_COL
                letter = openpyxl.utils.get_column_letter(col_index)
                cell = self.sheet[f"{letter}{row}"]
                cell.fill = PatternFill(fgColor= color["Fill Color"], fill_type= 'solid')


    def _set_format(self, col):
        col_index = self.cols.index(col) + self.START_COL
        for r in range(self.START_ROW + 1, self.sheet.max_row + 1):
            cell= self.sheet[openpyxl.utils.get_column_letter(col_index) + str(r)]
            cell.number_format = '@'
